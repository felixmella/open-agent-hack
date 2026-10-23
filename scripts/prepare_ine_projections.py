"""Download INE's official comuna population estimates and projections (2002-2035, base Censo 2017)
and aggregate them to one row per comuna and year.

Usage:  .venv\\Scripts\\python scripts\\prepare_ine_projections.py [--xlsx PATH]
Writes data/official/ine_projections_base2017_comunas.csv (cut, comuna, year, population).

The workbook has one row per comuna, sex and single year of age, with one column per year
("Poblacion 2002" ... "Poblacion 2035"); summing over sex and age gives the comuna total.
"""
import argparse
import csv
import io
import re
from collections import defaultdict
from pathlib import Path

import httpx
from openpyxl import load_workbook

SOURCE_URL = (
    "https://www.ine.gob.cl/docs/default-source/proyecciones-de-poblacion/cuadros-estadisticos/"
    "base-2017/estimaciones-y-proyecciones-2002-2035-comunas.xlsx"
)
OUTPUT = Path(__file__).resolve().parent.parent / "data" / "official" / "ine_projections_base2017_comunas.csv"
USER_AGENT = "FactSyncChile/0.1 (https://github.com/felixmella/open-agent-hack)"


def load_workbook_bytes(path):
    if path:
        return Path(path).read_bytes()
    response = httpx.get(SOURCE_URL, headers={"User-Agent": USER_AGENT}, timeout=180, follow_redirects=True)
    response.raise_for_status()
    return response.content


def aggregate(workbook_bytes):
    sheet = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    comuna_col, name_col = header.index("Comuna"), header.index("Nombre Comuna")
    year_cols = {i: int(re.search(r"(\d{4})", h).group(1)) for i, h in enumerate(header) if h and str(h).startswith("Poblacion")}
    totals, names = defaultdict(int), {}
    for row in rows:
        if row[comuna_col] is None:
            continue
        cut = str(int(row[comuna_col])).zfill(5)
        names[cut] = row[name_col]
        for i, year in year_cols.items():
            totals[cut, year] += int(row[i] or 0)
    return [(cut, names[cut], year, population) for (cut, year), population in sorted(totals.items())]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", help="use a local copy of the INE workbook instead of downloading it")
    args = parser.parse_args()
    records = aggregate(load_workbook_bytes(args.xlsx))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["cut", "comuna", "year", "population"])
        writer.writerows(records)
    comunas = len({r[0] for r in records})
    years = sorted({r[2] for r in records})
    print(f"wrote {OUTPUT} ({len(records)} rows: {comunas} comunas x {len(years)} years, {years[0]}-{years[-1]})")


if __name__ == "__main__":
    main()
